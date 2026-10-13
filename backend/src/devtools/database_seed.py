"""Validate Excel data, migrate once and seed in one PostgreSQL transaction."""

import hashlib
import json
import math
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from legal_checker.core.migrations import migrate
from legal_checker.core.postgres import connect_postgres
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from psycopg import sql
from psycopg.types.json import Jsonb

from devtools.database_schema import BY_NAME, FOREIGN_KEYS, GUIDE_SHEETS, TABLES, Table
from devtools.qa_schema import QA_GUIDES, QA_SHEETS, QA_TABLES


class SeedError(ValueError):
    pass


@dataclass
class WorkbookSeed:
    source_sha256: str
    tables: dict[str, list[dict[str, Any]]]
    qa_tables: dict[str, list[dict[str, Any]]] = field(default_factory=dict)

    @property
    def counts(self) -> dict[str, int]:
        return {name: len(rows) for name, rows in self.tables.items()}

    @property
    def qa_counts(self) -> dict[str, int]:
        return {name: len(rows) for name, rows in self.qa_tables.items()}


def convert_cell(value: Any, kind: str, epoch: datetime) -> Any:
    if value is None:
        return None
    if kind == "TEXT":
        return str(value)
    if kind == "JSONB":

        def invalid_constant(constant):
            raise ValueError(f"Non-finite JSON constant {constant}")

        def finite_float(raw):
            number = float(raw)
            if not math.isfinite(number):
                raise ValueError("Non-finite JSON number")
            return number

        return json.loads(str(value), parse_constant=invalid_constant, parse_float=finite_float)
    if kind == "BOOLEAN":
        if value in (True, 1, "1", "true", "TRUE"):
            return True
        if value in (False, 0, "0", "false", "FALSE"):
            return False
        raise ValueError("Expected boolean or 0/1; UNKNOWN cannot become False")
    if kind in ("INTEGER", "NUMERIC"):
        if isinstance(value, bool):
            raise ValueError("Boolean cannot become a number")
        try:
            number = Decimal(str(value))
        except InvalidOperation as error:
            raise ValueError("Expected a number") from error
        if not number.is_finite():
            raise ValueError("Expected finite number")
        if kind == "INTEGER":
            if number != number.to_integral_value():
                raise ValueError("Expected whole number")
            return int(number)
        return number
    if kind in ("DATE", "TIMESTAMPTZ"):
        if isinstance(value, (float, int)) and not isinstance(value, bool):
            if not math.isfinite(value):
                raise ValueError("Invalid Excel date")
            value = from_excel(value, epoch)
        if isinstance(value, str):
            # Excel contains leading apostrophes used to preserve ISO timestamps as text.
            value = datetime.fromisoformat(value.removeprefix("'").replace("Z", "+00:00"))
        if isinstance(value, date) and not isinstance(value, datetime):
            value = datetime.combine(value, datetime.min.time())
        if not isinstance(value, datetime):
            raise ValueError("Expected an ISO or Excel date")
        if kind == "DATE":
            return value.date()
        # Excel serial/naive audit dates have no timezone. Seed convention: UTC.
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value
    raise ValueError(f"Unsupported type {kind}")


def validate_references(seed: WorkbookSeed) -> None:
    for table in TABLES:
        rows = seed.tables[table.name]
        identifiers = [row[table.primary_key] for row in rows]
        if any(not key for key in identifiers) or len(set(identifiers)) != len(identifiers):
            raise SeedError(f"Empty or duplicate primary key in {table.name}")
    for source, column, target, key in FOREIGN_KEYS:
        identifiers = {row[key] for row in seed.tables[target]}
        for row in seed.tables[source]:
            if row[column] is not None and row[column] not in identifiers:
                raise SeedError(f"{source}.{column} references missing {target}.{key}")
    # Ordinary FKs cannot detect evidence borrowed from a different contract.
    cases = {r["ca_id"]: r["hop_dong_id"] for r in seed.tables["ca_danh_gia"]}
    requirements = {r["yeu_cau_ca_id"]: r for r in seed.tables["yeu_cau_ap_dung"]}
    chunks = {r["doan_id"]: r["hop_dong_id"] for r in seed.tables["doan_hop_dong"]}
    for row in seed.tables["ket_qua_truy_xuat"]:
        requirement = requirements[row["yeu_cau_ca_id"]]
        if chunks[row["doan_id"]] != cases[requirement["ca_id"]]:
            raise SeedError("Retrieval evidence belongs to a different contract")


def read_workbook(path: Path) -> WorkbookSeed:
    content = path.read_bytes()
    content_sha = hashlib.sha256(content).hexdigest()
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        unexpected = (
            set(workbook.sheetnames) - set(BY_NAME) - GUIDE_SHEETS - set(QA_SHEETS) - QA_GUIDES
        )
        missing = set(BY_NAME) - set(workbook.sheetnames)
        if missing or unexpected:
            raise SeedError(
                f"Workbook sheets differ: missing={sorted(missing)}, extra={sorted(unexpected)}"
            )
        data, qa_data = {}, {}
        sheet_tables = [(table.name, table) for table in TABLES] + list(QA_SHEETS.items())
        for sheet_name, table in sheet_tables:
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows()
            headers = tuple(cell.value for cell in next(rows))
            if headers != tuple(name for name, _ in table.columns):
                raise SeedError(f"Header mismatch in {table.name}")
            records = []
            for number, cells in enumerate(rows, 2):
                if all(cell.value is None for cell in cells):
                    continue
                if len(cells) > len(table.columns):
                    raise SeedError(f"Unexpected cells after header in {table.name} row {number}")
                record = {}
                # Write-only Excel sheets omit trailing blank cells in the XML.
                for index, (name, kind) in enumerate(table.columns):
                    cell = cells[index] if index < len(cells) else None
                    if cell is not None and cell.data_type in ("f", "e"):
                        raise SeedError(f"Formula/error in {table.name}!{cell.coordinate}")
                    try:
                        record[name] = convert_cell(
                            cell.value if cell is not None else None, kind, workbook.epoch
                        )
                    except (ValueError, TypeError, OverflowError) as error:
                        raise SeedError(f"{table.name} row {number}, {name}: {error}") from error
                records.append(record)
            (data if table.schema != "qa" else qa_data)[table.name] = records
        seed = WorkbookSeed(content_sha, data, qa_data)
        validate_references(seed)
        if qa_data:
            from devtools.qa_validation import validate_qa

            validate_qa(seed)
        # Inventory is optional for test workbooks; the original workbook includes it.
        if "BANG_VA_SO_DONG" in workbook.sheetnames:
            rows = workbook["BANG_VA_SO_DONG"].iter_rows(values_only=True)
            next(rows)
            inventory = {row[0]: row[1] for row in rows if row[0] in BY_NAME}
            if inventory != seed.counts:
                raise SeedError("Workbook inventory row counts differ from actual records")
        return seed
    finally:
        workbook.close()


def insert_table(connection, table: Table, rows: list[dict[str, Any]]) -> int:
    columns = [name for name, _ in table.columns]
    statement = sql.SQL("INSERT INTO {} ({}) VALUES ({}) ON CONFLICT ({}) DO NOTHING").format(
        sql.Identifier(table.schema, table.name),
        sql.SQL(", ").join(map(sql.Identifier, columns)),
        sql.SQL(", ").join(sql.Placeholder() for _ in columns),
        sql.Identifier(table.primary_key),
    )
    values = [
        tuple(
            Jsonb(row[name]) if kind == "JSONB" and row[name] is not None else row[name]
            for name, kind in table.columns
        )
        for row in rows
    ]
    with connection.cursor() as cursor:
        cursor.executemany(statement, values)
        return cursor.rowcount


def seed_database(database_url: str, seed: WorkbookSeed) -> dict:
    validate_references(seed)
    if seed.qa_tables:
        from devtools.qa_validation import validate_qa

        validate_qa(seed)
    with connect_postgres(database_url) as connection:
        # Serialize migration/seed across shells; all DDL and rows roll back together on error.
        connection.execute("SELECT pg_advisory_xact_lock(743249)")
        migrate(connection)
        prior = connection.execute("SELECT source_sha256 FROM dev_meta.seed_imports").fetchall()
        if prior:
            if any(row["source_sha256"] != seed.source_sha256 for row in prior):
                raise SeedError("Different workbook already seeded. Use a separate fresh local DB")
            return {
                "status": "already_seeded",
                "inserted": 0,
                "source_rows": sum(seed.counts.values()),
            }
        inserted = sum(insert_table(connection, table, seed.tables[table.name]) for table in TABLES)
        qa_inserted = sum(
            insert_table(connection, table, seed.qa_tables.get(table.name, []))
            for table in QA_TABLES
        )
        connection.execute("SET CONSTRAINTS ALL IMMEDIATE")
        connection.execute(
            "INSERT INTO dev_meta.seed_imports VALUES (%s, %s, now())",
            (
                seed.source_sha256,
                Jsonb({**seed.counts, **{"qa." + k: v for k, v in seed.qa_counts.items()}}),
            ),
        )
        return {
            "status": "seeded",
            "inserted": inserted,
            "skipped_existing": sum(seed.counts.values()) - inserted,
            "source_rows": sum(seed.counts.values()),
            "qa_inserted": qa_inserted,
        }


def database_counts(database_url: str) -> dict[str, int]:
    with connect_postgres(database_url) as connection:
        return {
            f"{table.schema}.{table.name}": connection.execute(
                sql.SQL("SELECT count(*) AS total FROM {}").format(
                    sql.Identifier(table.schema, table.name)
                )
            ).fetchone()["total"]
            for table in (*TABLES, *QA_TABLES)
            if table.schema != "qa"
            or connection.execute(
                "SELECT to_regclass(%s) AS name", ("qa." + table.name,)
            ).fetchone()["name"]
            is not None
        }
