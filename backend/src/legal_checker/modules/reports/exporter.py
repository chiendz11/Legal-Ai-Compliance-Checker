"""Export one exact C13 version to DOCX or PDF.

Both formats render the same block list, built only from the stored report: no statuses
are recomputed. Unpublished reports carry a visible DRAFT/REVIEW marking on every export.
"""

import io
import os
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol

from legal_checker.shared.contracts.report import FinalReport

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ReportExporter(Protocol):
    def export(self, report: FinalReport, format: str) -> bytes: ...


class ExportUnavailable(RuntimeError):
    """The renderer for this format is not installed/configured on this server."""


@dataclass
class Document:
    blocks: list[tuple] = field(default_factory=list)

    def heading(self, text: str, level: int = 1):
        self.blocks.append(("heading", level, text))

    def paragraph(self, text: str, *, emphasis: bool = False):
        self.blocks.append(("paragraph", text, emphasis))

    def table(self, headers: list[str], rows: list[list[str]]):
        self.blocks.append(("table", headers, rows))

    def quote(self, text: str, caption: str):
        self.blocks.append(("quote", text, caption))


def _v(value) -> str:
    if value is None:
        return "chưa có"
    if isinstance(value, bool):
        return "Có" if value else "Không"
    return str(value)


def report_document(report: FinalReport) -> Document:
    content = report.report_json
    case = content.case
    doc = Document()
    doc.heading(f"Báo cáo đánh giá hợp đồng {case.hop_dong_id}")
    if report.report_status != "PUBLISHED":
        doc.paragraph(
            f"{report.report_status} / review {report.review_status}: kết quả máy, chưa được "
            "chuyên viên xác nhận. Không dùng như kết luận pháp lý đã duyệt.",
            emphasis=True,
        )
    if content.audit.get("module_source") == "REFERENCE_REPLAY":
        doc.paragraph(
            "REFERENCE_REPLAY: dữ liệu tham chiếu QA được phát lại để kiểm thử tích hợp, "
            "không phải output của các module đánh giá.",
            emphasis=True,
        )
    context = content.audit.get("context", {})
    doc.table(
        ["Mục", "Giá trị"],
        [
            ["Mã báo cáo", report.bao_cao_id],
            ["Phiên bản", report.phien_ban_bao_cao],
            ["Ca", case.ca_id],
            ["Run", _v(context.get("run_id"))],
            ["Ngày đánh giá", _v(case.assessment_context.ngay_danh_gia)],
            [
                "Mốc áp dụng",
                f"{_v(case.assessment_context.moc_ap_dung)} "
                f"({_v(case.assessment_context.loai_moc_ap_dung)})",
            ],
            ["Rule", case.deployment.rule_version],
            ["Tri thức pháp lý", case.deployment.legal_knowledge_version],
            ["Snapshot tri thức", _v(context.get("knowledge_snapshot_sha256"))],
        ],
    )

    compliance = content.compliance
    doc.heading("Tổng quan", 2)
    requirement_counts = Counter(r.status for r in compliance.requirement_results)
    finding_counts = Counter(c.finding for c in compliance.criterion_results)
    citation_counts = Counter(f.status for f in content.citation_validation.findings)
    doc.table(
        ["Nhóm", "Số lượng theo trạng thái"],
        [
            ["Yêu cầu (C08)", ", ".join(f"{k}: {n}" for k, n in requirement_counts.items())],
            ["Tiêu chí (C07)", ", ".join(f"{k}: {n}" for k, n in finding_counts.items())],
            ["Viện dẫn (C09)", ", ".join(f"{k}: {n}" for k, n in citation_counts.items())],
            [
                "Rủi ro (C10)",
                f"COMPLIANCE: {len(content.risks.compliance)}, CITATION: "
                f"{len(content.risks.citation)}, VERIFICATION: {len(content.risks.verification)}",
            ],
            ["Khuyến nghị (C11)", str(len(content.recommendations))],
        ],
    )
    doc.paragraph("Không có điểm tuân thủ tổng: chỉ đếm theo trạng thái đã lưu.")

    doc.heading("Hồ sơ hợp đồng", 2)
    doc.table(
        ["Trường", "Giá trị"],
        [[k, _v(v)] for k, v in content.contract_profile.profile.model_dump(mode="json").items()],
    )

    doc.heading("Phạm vi pháp lý", 2)
    doc.table(
        ["Yêu cầu", "Phạm vi", "Phiên bản", "Áp dụng", "Mốc"],
        [
            [
                r.requirement_id,
                r.scope_status,
                f"{r.version_status} {_v(r.requirement_version_id)}",
                r.applicability_status,
                f"{_v(r.loai_moc)} {_v(r.moc_ap_dung)}",
            ]
            for r in content.applicable_legal_scope
        ],
    )

    doc.heading("Kết quả tuân thủ", 2)
    criteria_by_requirement: dict[str, list] = {}
    for criterion in compliance.criterion_results:
        criteria_by_requirement.setdefault(criterion.yeu_cau_ca_id, []).append(criterion)
    for requirement in compliance.requirement_results:
        version = requirement.requirement_snapshot_json.get("version", {})
        doc.heading(
            f"[{requirement.status}] {version.get('noi_dung') or requirement.yeu_cau_ca_id}", 3
        )
        doc.paragraph(
            f"{requirement.giai_thich} (đã đánh giá {requirement.so_tieu_chi_da_danh_gia}/"
            f"{requirement.so_tieu_chi_ap_dung}, chưa xác định {requirement.so_unknown})"
        )
        for criterion in criteria_by_requirement.get(requirement.yeu_cau_ca_id, []):
            text = criterion.tieu_chi_snapshot_json.get("criterion", {}).get("noi_dung")
            doc.paragraph(
                f"• [{criterion.finding}] {text or criterion.tieu_chi_id}: {criterion.giai_thich}"
                + (
                    f" — chưa chắc chắn: {criterion.uncertainty_reason}"
                    if criterion.uncertainty_reason
                    else ""
                )
            )
            for evidence in criterion.evidence:
                doc.quote(
                    evidence.trich_dan,
                    f"{evidence.doan_id}, ký tự {evidence.ky_tu_bat_dau}–{evidence.ky_tu_ket_thuc}",
                )
            if criterion.legal_basis_ids:
                doc.paragraph("Căn cứ: " + ", ".join(criterion.legal_basis_ids))

    doc.heading("Viện dẫn trong hợp đồng", 2)
    findings: dict[str, list] = {}
    for finding in content.citation_validation.findings:
        findings.setdefault(finding.vien_dan_id, []).append(finding)
    doc.table(
        ["Viện dẫn", "Tại ngày liên quan", "Hiện tại (tham khảo)", "Lý do"],
        [
            [
                citation.raw_text,
                "; ".join(
                    f"{f.status} @ {_v(f.relevant_date)}"
                    for f in findings.get(citation.vien_dan_id, [])
                )
                or "chưa xác minh",
                "; ".join(_v(f.current_status) for f in findings.get(citation.vien_dan_id, [])),
                "; ".join(f.reason_code for f in findings.get(citation.vien_dan_id, [])),
            ]
            for citation in content.citation_validation.citations
        ],
    )

    doc.heading("Rủi ro", 2)
    for title, risks in (
        ("Tuân thủ", content.risks.compliance),
        ("Viện dẫn", content.risks.citation),
        ("Cần xác minh", content.risks.verification),
    ):
        doc.heading(f"{title} ({len(risks)})", 3)
        if risks:
            doc.table(
                ["Mức độ", "Rủi ro", "Nguồn"],
                [
                    [
                        r.severity,
                        str(r.risk_snapshot.get("ten_rui_ro") or r.rui_ro_id),
                        r.source.ket_qua_tieu_chi_id or r.source.citation_finding_id or "",
                    ]
                    for r in risks
                ],
            )

    doc.heading("Khuyến nghị", 2)
    for recommendation in content.recommendations:
        doc.paragraph(
            f"• [{recommendation.uu_tien}] {recommendation.loai_hanh_dong}: "
            f"{recommendation.noi_dung} (review {recommendation.review_status})"
        )
    return doc


def _docx(document: Document) -> bytes:
    try:
        from docx import Document as Docx
        from docx.shared import Pt, RGBColor
    except ImportError as error:  # pragma: no cover - depends on installed extras
        raise ExportUnavailable("Install the backend 'export' extra for DOCX") from error
    out = Docx()
    out.styles["Normal"].font.size = Pt(10)
    for block in document.blocks:
        kind = block[0]
        if kind == "heading":
            out.add_heading(block[2], level=block[1])
        elif kind == "paragraph":
            run = out.add_paragraph().add_run(block[1])
            if block[2]:
                run.bold = True
                run.font.color.rgb = RGBColor(0x8A, 0x53, 0x00)
        elif kind == "quote":
            paragraph = out.add_paragraph(style="Intense Quote")
            paragraph.add_run(block[1])
            out.add_paragraph(block[2]).runs[0].font.size = Pt(8)
        elif kind == "table":
            headers, rows = block[1], block[2]
            table = out.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, headers, strict=True):
                cell.text = text
            for row in rows:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = text
    buffer = io.BytesIO()
    out.save(buffer)
    return buffer.getvalue()


FONT_CANDIDATES = (
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    ),
)


def _font() -> tuple[str, str]:
    configured = os.getenv("REPORT_PDF_FONT")
    if configured:
        return configured, os.getenv("REPORT_PDF_FONT_BOLD", configured)
    for regular, bold in FONT_CANDIDATES:
        if Path(regular).is_file():
            return regular, bold if Path(bold).is_file() else regular
    raise ExportUnavailable(
        "No Unicode TTF font for Vietnamese PDF; set REPORT_PDF_FONT to a .ttf file"
    )


# Word keeps Symbol/Wingdings glyphs as private-use code points (U+F0xx) that report fonts
# cannot draw; map the known ones to Unicode look-alikes, matching frontend/src/shared/text.ts.
_SYMBOL_FONT = {
    "\uf0b7": "\u2022",  # Symbol bullet
    "\uf0a7": "\u25aa",  # Wingdings small square
    "\uf06c": "\u25cf",  # Wingdings filled circle
    "\uf06e": "\u25a0",  # Wingdings filled square
    "\uf076": "\u2756",  # Wingdings diamond
    "\uf0d8": "\u27a2",  # Wingdings arrowhead
    "\uf0fc": "\u2713",  # Wingdings check mark
    "\uf0fe": "\u2611",  # Wingdings checked box
    "\uf02d": "-",  # Symbol minus
    "\uf02b": "+",  # Symbol plus
    "\u2028": "\n",  # line separator
    "\u2029": "\n",  # paragraph separator
}
_UNDRAWABLE = re.compile("[\ue000-\uf8ff\u2028\u2029]")


def _clean(value):
    if isinstance(value, str):
        # Unknown private-use glyphs become "?" rather than a missing-glyph box.
        return _UNDRAWABLE.sub(lambda m: _SYMBOL_FONT.get(m.group(), "?"), value)
    if isinstance(value, list):
        return [_clean(item) for item in value]
    return value


def _pdf(document: Document) -> bytes:
    try:
        from fpdf import FPDF
        from fpdf.fonts import FontFace
    except ImportError as error:  # pragma: no cover - depends on installed extras
        raise ExportUnavailable("Install the backend 'export' extra for PDF") from error
    regular, bold = _font()
    pdf = FPDF()
    pdf.add_font("Body", "", regular)
    pdf.add_font("Body", "B", bold)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    width = pdf.epw
    sizes = {1: 16, 2: 13, 3: 11}
    for block in document.blocks:
        kind = block[0]
        if kind == "heading":
            pdf.set_font("Body", "B", sizes.get(block[1], 11))
            pdf.ln(2)
            pdf.multi_cell(width, 7, block[2], new_x="LMARGIN", new_y="NEXT")
        elif kind == "paragraph":
            pdf.set_font("Body", "B" if block[2] else "", 9)
            if block[2]:
                pdf.set_text_color(138, 83, 0)
            pdf.multi_cell(width, 5, block[1], new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(0, 0, 0)
        elif kind == "quote":
            pdf.set_font("Body", "", 9)
            pdf.set_x(pdf.l_margin + 6)
            pdf.multi_cell(width - 6, 5, f"“{block[1]}”", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Body", "", 7)
            pdf.set_x(pdf.l_margin + 6)
            pdf.multi_cell(width - 6, 4, block[2], new_x="LMARGIN", new_y="NEXT")
        elif kind == "table":
            pdf.set_font("Body", "", 8)
            with pdf.table(
                headings_style=FontFace(emphasis="BOLD"), line_height=4.5, padding=1
            ) as table:
                table.row(block[1])
                for row in block[2]:
                    table.row(row)
            pdf.ln(2)
    return bytes(pdf.output())


def export_report(report: FinalReport, fmt: str) -> bytes:
    if fmt not in MEDIA_TYPES:
        raise ValueError(f"Unsupported export format {fmt}")
    document = report_document(report)
    document.blocks = [tuple(_clean(part) for part in block) for block in document.blocks]
    return _docx(document) if fmt == "docx" else _pdf(document)
