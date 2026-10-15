"""Small in-memory/workbook fixtures for seed tests; no dataset generation."""

import hashlib
import json
from datetime import date, datetime

from devtools.database_schema import BY_NAME, TABLES
from devtools.qa_schema import QA_TABLES
from openpyxl import Workbook


def native_row(table, **values):
    row = dict.fromkeys(k for k, _ in BY_NAME[table].columns)
    if set(values) - row.keys():
        raise ValueError("Unknown native columns: " + str(set(values) - row.keys()))
    return {**row, **values}


def artifact(qa, case_id, role, contract_name, record_id, value):
    """Encode a test payload using the frozen workbook's multipart representation."""
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    sha = hashlib.sha256(raw.encode()).hexdigest()
    parts = [raw[i : i + 28000] for i in range(0, len(raw), 28000)]
    for number, part in enumerate(parts, 1):
        qa["artifacts"].append(
            {
                "artifact_part_id": f"{case_id}/{role}/{contract_name}/{record_id}/{number}",
                "case_id": case_id,
                "role": role,
                "contract_name": contract_name,
                "record_id": record_id,
                "part_number": number,
                "part_count": len(parts),
                "payload_sha256": sha,
                "content": part,
            }
        )


def write_seed_workbook(seed, path):
    """Write minimal test sheets so the real reader can be tested without a release builder."""
    workbook = Workbook(write_only=True)
    for table in (*TABLES, *QA_TABLES):
        name = table.name if table.schema != "qa" else "QA_" + table.name.upper()
        sheet = workbook.create_sheet(name)
        sheet.append([name for name, _ in table.columns])
        rows = seed.tables if table.schema != "qa" else seed.qa_tables
        for row in rows.get(table.name, []):
            values = []
            for name, kind in table.columns:
                value = row.get(name)
                if value is not None:
                    if kind == "JSONB":
                        value = json.dumps(value, ensure_ascii=False)
                    elif isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    elif kind == "NUMERIC":
                        value = str(value)
                values.append(value)
            sheet.append(values)
    workbook.save(path)
