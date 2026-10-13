"""Workbook seed validation, independent of a running PostgreSQL server."""

import json
from datetime import date, datetime, timezone
from decimal import Decimal
from importlib.resources import files

import pytest
from devtools.database_migration import render_migration
from devtools.database_schema import BY_NAME, FOREIGN_KEYS, TABLES
from devtools.database_seed import SeedError, convert_cell, read_workbook
from openpyxl import Workbook
from openpyxl.utils.datetime import WINDOWS_EPOCH


def make_workbook(path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table in TABLES:
        sheet = workbook.create_sheet(table.name)
        sheet.append([name for name, _ in table.columns])
    workbook.save(path)
    return workbook


def test_migration_is_packaged_and_matches_frozen_contract():
    migration = files("legal_checker.core").joinpath("sql/001_workbook_v7_4_9.sql").read_text()
    assert migration == render_migration()
    assert len(TABLES) == 50
    assert sum(table.schema == "knowledge" for table in TABLES) == 22
    assert len(FOREIGN_KEYS) == 78


@pytest.mark.parametrize(
    ("value", "kind", "expected"),
    [
        ("CO", "TEXT", "CO"),
        ("UNKNOWN", "TEXT", "UNKNOWN"),
        ("0", "BOOLEAN", False),
        (True, "BOOLEAN", True),
        ("2.00", "INTEGER", 2),
        ("2200000000.01", "NUMERIC", Decimal("2200000000.01")),
        (
            '{"verified":false,"status":"UNKNOWN"}',
            "JSONB",
            {"verified": False, "status": "UNKNOWN"},
        ),
        (
            "'2026-10-06T14:56:29+00:00",
            "TIMESTAMPTZ",
            datetime(2026, 10, 6, 14, 56, 29, tzinfo=timezone.utc),
        ),
        (46305.5, "TIMESTAMPTZ", datetime(2026, 10, 10, 12, tzinfo=timezone.utc)),
        (datetime(2026, 10, 10), "DATE", date(2026, 10, 10)),
    ],
)
def test_excel_conversion(value, kind, expected):
    assert convert_cell(value, kind, WINDOWS_EPOCH) == expected


@pytest.mark.parametrize(
    ("value", "kind"),
    [
        ("UNKNOWN", "BOOLEAN"),
        ("CO", "BOOLEAN"),
        ("1.25", "INTEGER"),
        ("NaN", "NUMERIC"),
        ("abc", "NUMERIC"),
        ('{"score":1e400}', "JSONB"),
        ('{"score":NaN}', "JSONB"),
        ("not a date", "DATE"),
    ],
)
def test_conversion_rejects_invalid_data(value, kind):
    with pytest.raises(ValueError):
        convert_cell(value, kind, WINDOWS_EPOCH)


@pytest.mark.parametrize("problem", ["header", "duplicate", "foreign_key", "json", "formula"])
def test_reject_bad_workbook_before_database_writes(tmp_path, problem):
    path = tmp_path / "sample.xlsx"
    workbook = make_workbook(path)
    sheet = workbook["hop_dong"]
    if problem == "header":
        sheet["A1"] = "wrong_id"
    elif problem == "duplicate":
        sheet.append(["DOC-1"])
        sheet.append(["DOC-1"])
    elif problem == "foreign_key":
        workbook["ho_so_hop_dong"].append(["MISSING-DOC"])
    elif problem == "json":
        target = workbook["quy_tac_quyet_dinh"]
        row = {"quy_tac_id": "RULE-1", "dieu_kien_json": "{broken"}
        target.append([row.get(name) for name, _ in BY_NAME[target.title].columns])
    else:
        sheet.append(["=CONCAT(1,2)"])
    workbook.save(path)
    with pytest.raises(SeedError):
        read_workbook(path)


def test_workbook_preserves_json_and_provenance(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = make_workbook(path)
    workbook["hop_dong"].append(["DOC-1"])
    row = {
        "hop_dong_id": "DOC-1",
        "la_hop_dong_theo_mau": "UNKNOWN",
        "ben_ban_co_von_nuoc_ngoai": "CO",
        "gia_tri_hop_dong": 12345,
        "loai_du_lieu": "SYNTHETIC_DEMO",
    }
    workbook["ho_so_hop_dong"].append(
        [row.get(name) for name, _ in BY_NAME["ho_so_hop_dong"].columns]
    )
    row = {
        "quy_tac_tong_hop_id": "RULE-1",
        "dieu_kien_json": json.dumps({"trang_thai": "UNKNOWN"}, ensure_ascii=False),
    }
    workbook["quy_tac_tong_hop"].append(
        [row.get(name) for name, _ in BY_NAME["quy_tac_tong_hop"].columns]
    )
    workbook.save(path)
    seed = read_workbook(path)
    profile = seed.tables["ho_so_hop_dong"][0]
    assert profile["la_hop_dong_theo_mau"] == "UNKNOWN"
    assert profile["ben_ban_co_von_nuoc_ngoai"] == "CO"
    assert profile["loai_du_lieu"] == "SYNTHETIC_DEMO"
    assert seed.tables["quy_tac_tong_hop"][0]["dieu_kien_json"] == {"trang_thai": "UNKNOWN"}
