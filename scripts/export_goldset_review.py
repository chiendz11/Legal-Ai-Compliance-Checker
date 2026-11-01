"""Write the relevance labels to an Excel sheet for human review (column `nhan_sua`).

python -m scripts.export_goldset_review
python -m scripts.eval_retrieval_goldset --labels data/private/eval/goldset_review_DEV.xlsx
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
EVAL = ROOT / "data/private/eval"
GUIDE = [
    ("2", "Đoạn trả lời trực tiếp tiêu chí (đủ để người 3 kết luận)"),
    ("1", "Liên quan một phần / chỉ là tiêu đề, định nghĩa, ngữ cảnh"),
    ("0", "Không liên quan"),
    ("nhan_sua", "Để trống nếu đồng ý; ghi 0/1/2 nếu muốn sửa nhãn"),
]


def main() -> None:
    pool = json.loads((EVAL / "goldset_pool_DEV.json").read_text(encoding="utf-8"))
    labels = json.loads((EVAL / "goldset_labels_DEV.json").read_text(encoding="utf-8"))
    workbook = Workbook()
    guide = workbook.active
    guide.title = "Huong_dan"
    for row in [("Nhãn", "Ý nghĩa"), *GUIDE]:
        guide.append(row)
    guide.column_dimensions["B"].width = 70
    sheet = workbook.create_sheet("Nhan")
    columns = ("target", "ca_id", "tieu_chi_id", "tieu_chi", "doan_id", "dieu_khoan", "noi_dung")
    sheet.append([*columns, "nhan", "nhan_sua"])
    fills = {2: "D9EAD3", 1: "FFF2CC", 0: "FFFFFF"}
    for number, target in enumerate(pool):
        for candidate, grade in zip(target["candidates"], labels[str(number)], strict=True):
            sheet.append(
                [
                    number,
                    target["ca_id"],
                    target["tieu_chi_id"],
                    target["criterion"],
                    candidate["doan_id"],
                    candidate["path"],
                    " ".join(candidate["text"].split()),
                    grade,
                    None,
                ]
            )
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor=fills[grade])
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCDEFGHI", (7, 11, 30, 40, 24, 30, 90, 6, 9), strict=True):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = sheet.dimensions
    output = EVAL / "goldset_review_DEV.xlsx"
    workbook.save(output)
    print(f"Wrote {output} ({sheet.max_row - 1} labelled chunks)")


if __name__ == "__main__":
    main()
