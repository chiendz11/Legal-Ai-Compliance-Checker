"""Score retrieval variants on the pooled relevance labels (DEV, independent of reference C06).

Labels: 2 = the chunk answers the criterion, 1 = partial/context, 0 = not relevant.
Unjudged chunks count as 0; the pool holds the top-5 of every variant, so @1..@5 are fully
judged and @10 is a lower bound. Run from repository root after build_retrieval_goldset:

    python -m scripts.eval_retrieval_goldset
"""

import argparse
import json
import os
from collections import Counter, defaultdict
from pathlib import Path

from devtools.upstream_fixtures import load_artifact, load_context, load_document
from dotenv import load_dotenv
from legal_checker.modules.legal_knowledge.repository import Knowledge
from legal_checker.modules.retrieval.fact_evidence import matching_chunks
from legal_checker.modules.retrieval.repository import RetrievalDatasetRepository
from legal_checker.modules.retrieval.service import LexicalRetrievalService, chunk_index_version
from legal_checker.shared.contracts.temporal import ApplicableRequirement

from scripts.build_retrieval_goldset import _short_index, short_query

ROOT = Path(__file__).resolve().parents[1]
KS = (1, 3, 5, 10)
RRF_K = 60


def article(path: str | None) -> str:
    return "/".join((path or "").split("/")[:2])


def capped(order: list[str], paths: dict, cap: int) -> list[str]:
    seen: Counter = Counter()
    head, tail = [], []
    for doan_id in order:
        (head if seen[article(paths.get(doan_id))] < cap else tail).append(doan_id)
        seen[article(paths.get(doan_id))] += 1
    return head + tail


def fused(*orders: list[str]) -> list[str]:
    score: dict[str, float] = defaultdict(float)
    for order in orders:
        for rank, doan_id in enumerate(order, start=1):
            score[doan_id] += 1 / (RRF_K + rank)
    return sorted(score, key=lambda d: (-score[d], d))


def promoted(order: list[str], facts: list[str]) -> list[str]:
    return facts + [d for d in order if d not in facts]


def read_labels(path: Path, pool: list[dict]) -> dict[str, list[int]]:
    """JSON {target: [grade...]} or the review workbook, where a filled `nhan_sua` column
    overrides the annotator grade `nhan`."""
    if path.suffix != ".xlsx":
        return json.loads(path.read_text(encoding="utf-8"))
    from openpyxl import load_workbook

    rows = load_workbook(path, read_only=True)["Nhan"].iter_rows(values_only=True)
    header = next(rows)
    grades: dict[str, dict[str, int]] = defaultdict(dict)
    for values in rows:
        row = dict(zip(header, values, strict=False))
        override = row.get("nhan_sua")
        grade = row["nhan"] if override in (None, "") else override
        grades[str(row["target"])][row["doan_id"]] = int(grade)
    return {
        str(i): [grades[str(i)][c["doan_id"]] for c in target["candidates"]]
        for i, target in enumerate(pool)
    }


def main() -> None:
    load_dotenv(ROOT / ".env", override=False)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database-url", default=os.environ.get("DATABASE_URL"))
    parser.add_argument(
        "--pool", type=Path, default=ROOT / "data/private/eval/goldset_pool_DEV.json"
    )
    parser.add_argument(
        "--labels", type=Path, default=ROOT / "data/private/eval/goldset_labels_DEV.json"
    )
    args = parser.parse_args()
    db = args.database_url
    pool = json.loads(args.pool.read_text(encoding="utf-8"))
    labels = read_labels(args.labels, pool)
    knowledge, repository = Knowledge(db), RetrievalDatasetRepository(db)
    plain = LexicalRetrievalService()
    titled = LexicalRetrievalService(section_paths=repository.section_paths)
    selecting = {
        "bm25+path": titled,
        "bm25+path+fact": LexicalRetrievalService(
            section_paths=repository.section_paths, fact_evidence=repository.fact_evidence
        ),
    }

    hits: dict[str, Counter] = defaultdict(Counter)
    mrr: Counter = Counter()
    selection: dict[str, Counter] = defaultdict(Counter)
    targets = Counter()
    no_direct = []
    by_case = defaultdict(list)
    for number, target in enumerate(pool):
        by_case[target["ca_id"]].append((number, target))
    for ca_id, items in sorted(by_case.items()):
        document = load_document(db, ca_id)
        hop_dong_id = document.chunks[0].hop_dong_id
        paths = repository.section_paths(hop_dong_id)
        facts = repository.fact_evidence(hop_dong_id)
        context = load_context(db, ca_id, f"{ca_id}_B_GOLD")
        requirements = [
            ApplicableRequirement.model_validate(r)
            for r in load_artifact(db, ca_id, "EXPECTED", "C05")
        ]
        outputs = {
            name: service.retrieve(context, requirements, document.chunks, knowledge)
            for name, service in selecting.items()
        }
        version_plain = chunk_index_version(document.chunks, "PLAIN")
        plain.indexes["bm25"].index(hop_dong_id, document.chunks, version_plain)
        version_titled, chunk_titled = _short_index(titled, document.chunks)
        for number, target in items:
            grade = {
                c["doan_id"]: labels[str(number)][i] for i, c in enumerate(target["candidates"])
            }
            query = target["criterion"]

            def search(service, text, version, contract=hop_dong_id):
                hits = service.indexes["bm25"].search(contract, text, 30, version)
                return [h.doan_id for h in hits]

            base = search(plain, query, version_plain)
            path = search(titled, query, version_titled)
            short = search(titled, short_query(query), version_titled)
            fact_ids = [d for d, _ in matching_chunks(query, facts, limit=3) if d in chunk_titled]
            variants = {
                "bm25": base,
                "bm25+path": path,
                "bm25+path+cap3": capped(path, paths, 3),
                "bm25+path+short": fused(path, short),
                "bm25+path+fact": promoted(path, fact_ids),
                "bm25+path+fact+short": promoted(fused(path, short), fact_ids),
                "bm25+path+fact+cap3": promoted(capped(path, paths, 3), fact_ids),
            }
            has_direct = 2 in grade.values()
            has_any = any(g >= 1 for g in grade.values())
            targets["direct"] += has_direct
            targets["any"] += has_any
            if not has_direct:
                no_direct.append(f"{target['tieu_chi_id']} ({ca_id})")
            for name, order in variants.items():
                top = order[:10]
                for k in KS:
                    hits[name][f"strict@{k}"] += has_direct and any(
                        grade.get(d, 0) == 2 for d in top[:k]
                    )
                    hits[name][f"lenient@{k}"] += has_any and any(
                        grade.get(d, 0) >= 1 for d in top[:k]
                    )
                if has_direct:
                    first = next((i for i, d in enumerate(top, 1) if grade.get(d, 0) == 2), None)
                    mrr[name] += 1 / first if first else 0
            for name, rows in outputs.items():
                chosen = [
                    r.doan_id
                    for r in rows
                    if r.duoc_chon
                    and r.tieu_chi_id == target["tieu_chi_id"]
                    and r.yeu_cau_ca_id == target["yeu_cau_ca_id"]
                ]
                counts = selection[name]
                counts["selected"] += len(chosen)
                counts["relevant"] += sum(grade.get(d, 0) >= 1 for d in chosen)
                counts["direct"] += sum(grade.get(d, 0) == 2 for d in chosen)
                counts["targets_with_direct_selected"] += has_direct and any(
                    grade.get(d, 0) == 2 for d in chosen
                )
                counts["empty_selection"] += not chosen
                counts["empty_when_no_direct"] += (not chosen) and not has_direct

    n_direct, n_any = targets["direct"], targets["any"]
    print(
        f"{len(pool)} targets: {n_direct} with a direct (2) chunk, {n_any} with any relevant chunk"
    )
    print(f"No direct evidence in pool: {len(no_direct)} -> {', '.join(no_direct)}\n")
    header = (
        "variant".ljust(24)
        + "".join(f"strict@{k:<3}" for k in KS)
        + "  MRR   "
        + "".join(f"lenient@{k:<2}" for k in (1, 3))
    )
    print(header)
    for name in hits:
        row = "".join(f"{hits[name][f'strict@{k}'] / n_direct:>9.1%} " for k in KS)
        row += f" {mrr[name] / n_direct:5.3f} "
        row += "".join(f"{hits[name][f'lenient@{k}'] / n_any:>10.1%} " for k in (1, 3))
        print(name.ljust(24) + row)
    print("\nduoc_chon quality (real service output):")
    for name, c in selection.items():
        share = max(c["selected"], 1)
        print(
            f"  {name:16} selected {c['selected']:3} | relevant {c['relevant'] / share:.1%}"
            f" | direct {c['direct'] / share:.1%} | targets whose selection has a direct chunk"
            f" {c['targets_with_direct_selected']}/{n_direct} | empty selections"
            f" {c['empty_selection']} ({c['empty_when_no_direct']} without direct evidence)"
        )


if __name__ == "__main__":
    main()
