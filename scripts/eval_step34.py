"""Compare step 3–4 implementations with seeded reference outputs on DEV cases.

Reference outputs are SILVER references, not legal gold: report every mismatch for review
instead of tuning the resolver until the numbers agree. Run from repository root.
"""

import argparse
import json
import os
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from devtools.upstream_fixtures import list_cases, load_artifact, load_context, load_document
from dotenv import load_dotenv
from legal_checker.modules.applicable_law.scope_resolver import CatalogScopeResolver, get_field
from legal_checker.modules.applicable_law.temporal_resolver import CatalogTemporalResolver
from legal_checker.modules.citation_validation.service import CatalogCitationValidator
from legal_checker.modules.legal_knowledge.repository import Knowledge
from legal_checker.modules.retrieval.repository import RetrievalDatasetRepository
from legal_checker.modules.retrieval.service import (
    LexicalRetrievalService,
    semantic_retrieval_service,
)
from legal_checker.shared.contracts.citation import CitationValidation
from legal_checker.shared.contracts.handoffs import LegalOutput
from legal_checker.shared.contracts.retrieval import RetrievalEvidence
from legal_checker.shared.contracts.scope import ScopeResolution
from legal_checker.shared.contracts.temporal import ApplicableRequirement
from legal_checker.shared.validation import validate_context, validate_requirements
from legal_checker.workflows.handoffs import validate_legal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
STATUSES = ("MATCH", "NO_MATCH", "UNKNOWN")
APPLICABILITY = ("CONFIRMED", "NOT_APPLICABLE", "UNKNOWN")
TEMPORAL_FIELDS = (
    "applicability_status",
    "version_status",
    "requirement_version_id",
    "loai_moc",
    "moc_ap_dung",
    "scope_status",
)
# C02 values the scope catalog conditions on, shown next to each comparison row.
PROFILE_COLUMNS = (
    "loai_hop_dong_id",
    "che_do_giao_dich",
    "tinh_trang_tai_san",
    "la_nha_o_xa_hoi",
    "muc_dich_giao_dich",
    "la_hop_dong_theo_mau",
    "ben_ban_la_chu_dau_tu",
    "la_nha_xay_moi",
    "BEN_BAN.la_chu_the_kinh_doanh",
    "ngay_giao_ket",
)
COMPARE_COLUMNS = (
    "yeu_cau_ca_id",
    "ca_id",
    "requirement_id",
    "khop",
    "ref_scope_status",
    "our_scope_status",
    "ref_nhom_phien_ban",
    "our_nhom_phien_ban",
    "ref_scope_reason",
    "our_scope_reason",
    *PROFILE_COLUMNS,
)
MISMATCH_FILL = PatternFill("solid", fgColor="FDE2E1")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")


def write_sheet(sheet, columns, rows, highlight=None) -> None:
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append([row.get(c) for c in columns])
        if highlight and highlight(row):
            for cell in sheet[sheet.max_row]:
                cell.fill = MISMATCH_FILL
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, start=1):
        width = max(len(str(column)), *(len(str(r.get(column) or "")) for r in rows[:200]), 8)
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 60)


def eval_scope(database_url: str, split: str, limit: int | None, output: Path) -> None:
    knowledge = Knowledge(database_url)
    cases = list_cases(database_url, split)[:limit]
    matrix: dict[str, Counter] = defaultdict(Counter)
    shape_errors = []
    compared = []
    resolver = None
    for ca_id in cases:
        reference = {
            row["yeu_cau_ca_id"]: ScopeResolution.model_validate(row)
            for row in load_artifact(database_url, ca_id, "EXPECTED", "C04")
        }
        if resolver is None:
            # Same supported catalog as the reference adapter (10 families).
            resolver = CatalogScopeResolver({r.requirement_id for r in reference.values()})
        context = load_context(database_url, ca_id, run_id=f"{ca_id}_B_DEV_EVAL")
        profile = load_document(database_url, ca_id).contract_profile
        profile_values = {name: get_field(profile, name).value for name in PROFILE_COLUMNS}
        results = resolver.resolve(context, profile, knowledge)
        ours = {r.yeu_cau_ca_id: r for r in results}
        if ours.keys() != reference.keys():
            shape_errors.append((ca_id, sorted(ours.keys() ^ reference.keys())))
        for key, ref in reference.items():
            mine = ours.get(key)
            if mine is None:
                continue
            # Round-trip through JSON exactly as a downstream reader would.
            ScopeResolution.model_validate(json.loads(mine.model_dump_json()))
            for name in ("profile_snapshot_json", "dieu_kien_snapshot_json"):
                missing = {"reference_scope_reason", "synthetic"}
                if not (ref.model_dump()[name].keys() - missing) <= mine.model_dump()[name].keys():
                    shape_errors.append((key, name))
            matrix[ref.requirement_id][(ref.scope_status, mine.scope_status)] += 1
            compared.append(
                {
                    "yeu_cau_ca_id": key,
                    "ca_id": ca_id,
                    "requirement_id": ref.requirement_id,
                    "khop": "CO" if ref.scope_status == mine.scope_status else "KHONG",
                    "ref_scope_status": ref.scope_status,
                    "our_scope_status": mine.scope_status,
                    "ref_nhom_phien_ban": ref.nhom_phien_ban,
                    "our_nhom_phien_ban": mine.nhom_phien_ban,
                    "ref_scope_reason": ref.scope_reason,
                    "our_scope_reason": mine.scope_reason,
                    **{
                        name: str(value) if value is not None else None
                        for name, value in profile_values.items()
                    },
                }
            )

    total = Counter()
    summary = []
    print(f"Scope on {len(cases)} {split} cases (rows: reference, columns: ours)")
    for requirement_id in sorted(matrix):
        counts = matrix[requirement_id]
        total.update(counts)
        agree = sum(counts[(s, s)] for s in STATUSES)
        print(f"\n{requirement_id}  agree {agree}/{sum(counts.values())}")
        print("  ref\\ours " + "".join(f"{s:>10}" for s in STATUSES))
        for ref in STATUSES:
            print(f"  {ref:9}" + "".join(f"{counts[(ref, s)]:>10}" for s in STATUSES))
            for ours_status in STATUSES:
                if counts[(ref, ours_status)]:
                    summary.append(
                        {
                            "requirement_id": requirement_id,
                            "ref_scope_status": ref,
                            "our_scope_status": ours_status,
                            "so_dong": counts[(ref, ours_status)],
                            "khop": "CO" if ref == ours_status else "KHONG",
                        }
                    )
    agree = sum(total[(s, s)] for s in STATUSES)
    print(f"\nOverall agreement {agree}/{sum(total.values())}")
    print(f"Shape errors: {len(shape_errors)} {shape_errors[:5]}")

    workbook = Workbook()
    mismatch = lambda row: row["khop"] == "KHONG"  # noqa: E731
    write_sheet(
        workbook.active,
        ("requirement_id", "ref_scope_status", "our_scope_status", "so_dong", "khop"),
        summary,
        mismatch,
    )
    workbook.active.title = "Tong_hop"
    write_sheet(
        workbook.create_sheet("Lech"), COMPARE_COLUMNS, [r for r in compared if mismatch(r)]
    )
    write_sheet(workbook.create_sheet("So_sanh_tat_ca"), COMPARE_COLUMNS, compared, mismatch)
    save_workbook(workbook, output)


def save_workbook(workbook: Workbook, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output)
    except PermissionError:
        # Excel locks an open workbook on Windows; keep the result under a new name.
        output = output.with_stem(f"{output.stem}_{datetime.now():%Y%m%d_%H%M%S}")
        workbook.save(output)
        print("Previous comparison is open in Excel; close it to reuse the default name.")
    print(f"Side-by-side comparison written to {output}")


def eval_temporal(database_url: str, split: str, limit: int | None, output: Path) -> None:
    """Run our C04 into our C05 (the real pipeline) and compare with reference C05."""
    knowledge = Knowledge(database_url)
    cases = list_cases(database_url, split)[:limit]
    temporal = CatalogTemporalResolver()
    matrix: dict[str, Counter] = defaultdict(Counter)
    field_agree: Counter = Counter()
    errors = []
    compared = []
    scope_resolver = None
    for ca_id in cases:
        reference = {
            row["yeu_cau_ca_id"]: ApplicableRequirement.model_validate(row)
            for row in load_artifact(database_url, ca_id, "EXPECTED", "C05")
        }
        if scope_resolver is None:
            scope_resolver = CatalogScopeResolver({r.requirement_id for r in reference.values()})
        context = load_context(database_url, ca_id, run_id=f"{ca_id}_B_DEV_EVAL")
        document = load_document(database_url, ca_id)
        scope = scope_resolver.resolve(context, document.contract_profile, knowledge)
        results = temporal.resolve(
            context, document.contract_profile, [s.model_copy(deep=True) for s in scope], knowledge
        )
        # Same handoff checks as DraftAssessmentWorkflow: checklist and scope fields unchanged.
        if [r.yeu_cau_ca_id for r in results] != [s.yeu_cau_ca_id for s in scope] or any(
            (r.requirement_id, r.scope_status) != (s.requirement_id, s.scope_status)
            for r, s in zip(results, scope, strict=True)
        ):
            errors.append((ca_id, "temporal changed scope-owned fields"))
        ours = {r.yeu_cau_ca_id: r for r in results}
        if ours.keys() != reference.keys():
            errors.append((ca_id, sorted(ours.keys() ^ reference.keys())))
        for key, ref in reference.items():
            mine = ours.get(key)
            if mine is None:
                continue
            ApplicableRequirement.model_validate(json.loads(mine.model_dump_json()))
            for name in ("nguon_snapshot_json", "reason"):
                if not ref.model_dump()[name].keys() <= mine.model_dump()[name].keys():
                    errors.append((key, f"{name} misses reference keys"))
            same = {field: getattr(ref, field) == getattr(mine, field) for field in TEMPORAL_FIELDS}
            field_agree.update(f for f, ok in same.items() if ok)
            matrix[ref.requirement_id][(ref.applicability_status, mine.applicability_status)] += 1
            compared.append(
                {
                    "yeu_cau_ca_id": key,
                    "ca_id": ca_id,
                    "requirement_id": ref.requirement_id,
                    "khop": "CO" if all(same.values()) else "KHONG",
                    "lech_o": ", ".join(f for f, ok in same.items() if not ok),
                    **{f"ref_{f}": str(getattr(ref, f)) for f in TEMPORAL_FIELDS},
                    **{f"our_{f}": str(getattr(mine, f)) for f in TEMPORAL_FIELDS},
                    "ref_reason": ref.reason.get("summary"),
                    "our_reason": mine.reason.get("summary"),
                    "our_scope_reason": mine.reason.get("scope"),
                    "ngay_giao_ket": str(document.contract_profile.profile.ngay_giao_ket),
                }
            )

    rows = len(compared)
    print(f"Temporal on {len(cases)} {split} cases, {rows} rows (our C04 -> our C05)")
    print("Field agreement: " + ", ".join(f"{f} {field_agree[f]}/{rows}" for f in TEMPORAL_FIELDS))
    summary = []
    for requirement_id in sorted(matrix):
        counts = matrix[requirement_id]
        agree = sum(counts[(s, s)] for s in APPLICABILITY)
        print(f"\n{requirement_id}  applicability agree {agree}/{sum(counts.values())}")
        print("  ref\\ours      " + "".join(f"{s:>16}" for s in APPLICABILITY))
        for ref in APPLICABILITY:
            print(f"  {ref:14}" + "".join(f"{counts[(ref, s)]:>16}" for s in APPLICABILITY))
            summary += [
                {
                    "requirement_id": requirement_id,
                    "ref_applicability_status": ref,
                    "our_applicability_status": ours_status,
                    "so_dong": counts[(ref, ours_status)],
                    "khop": "CO" if ref == ours_status else "KHONG",
                }
                for ours_status in APPLICABILITY
                if counts[(ref, ours_status)]
            ]
    full = sum(r["khop"] == "CO" for r in compared)
    print(f"\nRows equal on every compared field: {full}/{rows}")
    print(f"Handoff/shape errors: {len(errors)} {errors[:5]}")

    columns = (
        "yeu_cau_ca_id",
        "ca_id",
        "requirement_id",
        "khop",
        "lech_o",
        *(f"{side}_{f}" for f in TEMPORAL_FIELDS for side in ("ref", "our")),
        "ref_reason",
        "our_reason",
        "our_scope_reason",
        "ngay_giao_ket",
    )
    workbook = Workbook()
    mismatch = lambda row: row["khop"] == "KHONG"  # noqa: E731
    write_sheet(
        workbook.active,
        (
            "requirement_id",
            "ref_applicability_status",
            "our_applicability_status",
            "so_dong",
            "khop",
        ),
        summary,
        mismatch,
    )
    workbook.active.title = "Tong_hop"
    write_sheet(workbook.create_sheet("Lech"), columns, [r for r in compared if mismatch(r)])
    write_sheet(workbook.create_sheet("So_sanh_tat_ca"), columns, compared, mismatch)
    save_workbook(workbook, output)


def eval_retrieval(
    database_url: str,
    split: str,
    limit: int | None,
    output: Path,
    mode: str = "bm25",
    embedding_model: str | None = None,
) -> None:
    """Reference C05 in, our C06 out: isolates retrieval from scope/temporal differences.

    Reference selections are lexical spans, not relevance gold; Recall@K says how often our
    ranking surfaces the chunks the reference selected.
    """
    knowledge = Knowledge(database_url)
    cases = list_cases(database_url, split)[:limit]
    repository = RetrievalDatasetRepository(database_url)
    sources = {"section_paths": repository.section_paths, "fact_evidence": repository.fact_evidence}
    if mode == "bm25":
        service = LexicalRetrievalService(**sources)
    else:
        service = semantic_retrieval_service(
            **sources,
            embedding_model=embedding_model,
            cache_dir=ROOT / "data/private/embedding_cache",
            rerank=mode in ("rerank", "bm25_rerank"),
            dense=mode in ("hybrid", "rerank"),
        )
    print(f"Retrieval mode: {mode} ({service.retrieval_version})")
    ks = (1, 3, 5, 10)
    hits_at: Counter = Counter()
    targets = 0
    selected_pairs = {"ref": set(), "ours": set()}
    coverage: Counter = Counter()
    errors = []
    compared = []
    for ca_id in cases:
        context = load_context(database_url, ca_id, run_id=f"{ca_id}_B_DEV_EVAL")
        document = load_document(database_url, ca_id)
        chunk_text = {c.doan_id: c.noi_dung for c in document.chunks}
        requirements = [
            ApplicableRequirement.model_validate(row)
            for row in load_artifact(database_url, ca_id, "EXPECTED", "C05")
        ]
        reference = [
            RetrievalEvidence.model_validate(row)
            for row in load_artifact(database_url, ca_id, "EXPECTED", "C06")
        ]
        ours = service.retrieve(context, requirements, document.chunks, knowledge)
        for row in ours:
            RetrievalEvidence.model_validate(json.loads(row.model_dump_json()))
            if row.doan_id not in chunk_text:
                errors.append((row.truy_xuat_id, "chunk of another contract"))
        if len({r.truy_xuat_id for r in ours}) != len(ours):
            errors.append((ca_id, "duplicate truy_xuat_id"))

        def by_criterion(rows):
            grouped = defaultdict(list)
            for r in sorted(rows, key=lambda r: r.thu_hang):
                grouped[(r.yeu_cau_ca_id, r.tieu_chi_id)].append(r)
            return grouped

        ref_rows, our_rows = by_criterion(reference), by_criterion(ours)
        coverage["ref_criteria"] += len(ref_rows)
        coverage["our_criteria"] += len(our_rows)
        coverage["both"] += len(ref_rows.keys() & our_rows.keys())
        for key in sorted(ref_rows.keys() | our_rows.keys()):
            ref_list, our_list = ref_rows.get(key, []), our_rows.get(key, [])
            ref_selected = [r.doan_id for r in ref_list if r.duoc_chon]
            ranking = [r.doan_id for r in our_list]
            selected_pairs["ref"].update((key, d) for d in ref_selected)
            selected_pairs["ours"].update((key, r.doan_id) for r in our_list if r.duoc_chon)
            found = None
            if ref_selected:
                targets += 1
                for k in ks:
                    hits_at[k] += any(d in ranking[:k] for d in ref_selected)
                found = next((i + 1 for i, d in enumerate(ranking) if d in ref_selected), None)
            top = our_list[0] if our_list else None
            compared.append(
                {
                    "yeu_cau_ca_id": key[0],
                    "tieu_chi_id": key[1],
                    "query_text": (ref_list or our_list)[0].query_text,
                    "ref_selected": ", ".join(ref_selected),
                    "ref_selected_text": " || ".join(chunk_text[d][:200] for d in ref_selected),
                    "our_selected": ", ".join(r.doan_id for r in our_list if r.duoc_chon),
                    "our_top1_text": chunk_text[top.doan_id][:200] if top else None,
                    "our_top1_score": top.scores.get("bm25") if top else None,
                    "rank_of_ref_selected": found,
                    "khop": (
                        None
                        if not ref_selected
                        else ("CO" if found is not None and found <= 3 else "KHONG")
                    ),
                }
            )

    print(f"Retrieval on {len(cases)} {split} cases (reference C05 -> our C06)")
    print(
        f"Criteria with rows: reference {coverage['ref_criteria']}, ours "
        f"{coverage['our_criteria']}, both {coverage['both']}"
    )
    print(f"Criteria with reference-selected chunks: {targets}")
    for k in ks:
        print(f"  Recall@{k}: {hits_at[k]}/{targets} = {hits_at[k] / max(targets, 1):.1%}")
    both = selected_pairs["ref"] & selected_pairs["ours"]
    print(
        f"Selection (duoc_chon) pairs: ref {len(selected_pairs['ref'])}, ours "
        f"{len(selected_pairs['ours'])}, overlap {len(both)} "
        f"(precision {len(both) / max(len(selected_pairs['ours']), 1):.1%}, "
        f"recall {len(both) / max(len(selected_pairs['ref']), 1):.1%})"
    )
    print(f"Shape/ownership errors: {len(errors)} {errors[:5]}")

    columns = (
        "yeu_cau_ca_id",
        "tieu_chi_id",
        "khop",
        "rank_of_ref_selected",
        "query_text",
        "ref_selected",
        "ref_selected_text",
        "our_selected",
        "our_top1_score",
        "our_top1_text",
    )
    workbook = Workbook()
    mismatch = lambda row: row["khop"] == "KHONG"  # noqa: E731
    write_sheet(workbook.active, columns, [r for r in compared if r["khop"]], mismatch)
    workbook.active.title = "Co_tham_chieu_chon"
    write_sheet(workbook.create_sheet("Tat_ca_tieu_chi"), columns, compared, mismatch)
    save_workbook(workbook, output)


def eval_citation(database_url: str, split: str, limit: int | None, output: Path) -> None:
    """Our C09 on DEV. The reference C09 is all UNKNOWN (no gold), so this checks format
    against it and lists every finding for manual review, INVALID/STALE first."""
    knowledge = Knowledge(database_url)
    cases = list_cases(database_url, split)[:limit]
    validator = CatalogCitationValidator()
    statuses: Counter = Counter()
    reasons: Counter = Counter()
    verified = 0
    errors = []
    rows = []
    for ca_id in cases:
        context = load_context(database_url, ca_id, run_id=f"{ca_id}_B_DEV_EVAL")
        document = load_document(database_url, ca_id)
        reference = {
            row["vien_dan_id"]: CitationValidation.model_validate(row)
            for row in load_artifact(database_url, ca_id, "EXPECTED", "C09")
        }
        findings = validator.validate(
            context, document.contract_profile, document.citations, knowledge
        )
        raw = {c.vien_dan_id: c for c in document.citations}
        if {f.vien_dan_id for f in findings} != reference.keys():
            errors.append((ca_id, "citation set differs from reference C09"))
        for finding in findings:
            CitationValidation.model_validate(json.loads(finding.model_dump_json()))
            ref = reference.get(finding.vien_dan_id)
            if ref and ref.finding_id != finding.finding_id:
                errors.append((finding.vien_dan_id, "finding_id differs from reference"))
            if ref and not (
                ref.source_snapshot_json.keys() - {"synthetic"}
                <= finding.source_snapshot_json.keys()
            ):
                errors.append((finding.vien_dan_id, "source_snapshot_json misses keys"))
            statuses[finding.status] += 1
            reasons[(finding.status, finding.reason_code)] += 1
            verified += finding.finding_verified
            citation = raw[finding.vien_dan_id]
            rows.append(
                {
                    "finding_id": finding.finding_id,
                    "ca_id": ca_id,
                    "status": finding.status,
                    "reason_code": finding.reason_code,
                    "domain_status": finding.domain_status,
                    "finding_verified": finding.finding_verified,
                    "observed_number": citation.normalized.so_hieu,
                    "observed_title": citation.normalized.ten_van_ban,
                    "van_ban_id": finding.normalized_van_ban_id,
                    "relevant_date": str(finding.relevant_date),
                    "current_status": finding.current_status,
                    "replacement_van_ban_id": finding.replacement_van_ban_id,
                    "ly_do": finding.ly_do,
                    "ref_status": ref.status if ref else None,
                    "raw_text": " ".join(citation.raw_text.split())[:300],
                }
            )

    total = sum(statuses.values())
    print(f"Citation on {len(cases)} {split} cases: {total} findings")
    for status in ("VALID", "STALE", "INVALID", "UNKNOWN"):
        print(f"  {status:8} {statuses[status]:6} ({statuses[status] / max(total, 1):.1%})")
    for (status, code), count in sorted(reasons.items()):
        print(f"    {status:8} {code:40} {count}")
    print(f"finding_verified=true: {verified}")
    print(f"Format/ID errors vs reference C09: {len(errors)} {errors[:5]}")

    order = {"INVALID": 0, "STALE": 1, "VALID": 2, "UNKNOWN": 3}
    rows.sort(key=lambda r: (order[r["status"]], r["reason_code"], r["finding_id"]))
    workbook = Workbook()
    columns = tuple(rows[0]) if rows else ("finding_id",)
    flagged = lambda row: row["status"] in ("INVALID", "STALE")  # noqa: E731
    write_sheet(workbook.active, columns, [r for r in rows if flagged(r)], flagged)
    workbook.active.title = "INVALID_STALE"
    write_sheet(workbook.create_sheet("Tat_ca"), columns, rows, flagged)
    save_workbook(workbook, output)


def eval_handoff(database_url: str, split: str, limit: int | None, pin_local: bool) -> None:
    """Run steps 3–4 together through the workflow's own handoff validators.

    These are the checks DraftAssessmentWorkflow applies when the real step 1–2 and 5–7
    implementations are plugged in, so a failure here is a failure at integration time.
    """
    knowledge = Knowledge(database_url)
    cases = list_cases(database_url, split)[:limit]
    families = {
        r["requirement_id"] for r in load_artifact(database_url, cases[0], "EXPECTED", "C04")
    }
    scope_resolver, temporal = CatalogScopeResolver(families), CatalogTemporalResolver()
    repository = RetrievalDatasetRepository(database_url)
    retrieval = LexicalRetrievalService(
        section_paths=repository.section_paths, fact_evidence=repository.fact_evidence
    )
    citations = CatalogCitationValidator()
    failures: Counter = Counter()
    for ca_id in cases:
        context = load_context(database_url, ca_id, run_id=f"{ca_id}_B_DEV_001")
        if pin_local:
            context = context.model_copy(update={"knowledge_snapshot_sha256": knowledge.sha256})
        document = load_document(database_url, ca_id)
        try:
            validate_context(context, knowledge)
            scope = scope_resolver.resolve(context, document.contract_profile, knowledge)
            requirements = temporal.resolve(
                context,
                document.contract_profile,
                [s.model_copy(deep=True) for s in scope],
                knowledge,
            )
            validate_requirements(context, requirements, knowledge)
            retrieval.retrieve(context, requirements, document.chunks, knowledge)
            findings = citations.validate(
                context, document.contract_profile, document.citations, knowledge
            )
            validate_legal(
                context,
                document,
                LegalOutput(scope=scope, requirements=requirements, citations=findings),
                knowledge,
            )
        except ValueError as error:
            failures[str(error)] += 1
    passed = len(cases) - sum(failures.values())
    print(f"Workflow handoff checks on {len(cases)} {split} cases: {passed} passed")
    if pin_local:
        print("  (context pinned to the locally captured knowledge hash)")
    for message, count in failures.most_common():
        print(f"  {count:4} {message}")


def main() -> None:
    load_dotenv(ROOT / ".env", override=False)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database-url", default=os.environ.get("DATABASE_URL"))
    parser.add_argument("--split", default="DEV", choices=["DEV", "VALIDATION"])
    parser.add_argument("--limit", type=int, help="Only the first N cases")
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("scope", help="Step 3A.1 C04 against reference C04")
    commands.add_parser("temporal", help="Step 3A.2 C05 (from our C04) against reference C05")
    retrieval = commands.add_parser(
        "retrieval", help="Step 4 C06 (from reference C05) against reference C06"
    )
    retrieval.add_argument(
        "--mode",
        choices=["bm25", "bm25_rerank", "hybrid", "rerank"],
        default="bm25",
        help="bm25; bm25_rerank; hybrid = BM25+embedding (RRF); rerank = hybrid + cross-encoder",
    )
    retrieval.add_argument("--embedding-model", help="Sentence-transformers model name")
    commands.add_parser("citation", help="Step 3B C09 format check and review workbook")
    handoff = commands.add_parser("handoff", help="Steps 3–4 through workflow handoff validators")
    handoff.add_argument(
        "--pin-local-knowledge",
        action="store_true",
        help="Use the locally captured knowledge hash when it differs from the manifest pin",
    )
    args = parser.parse_args()
    if not args.database_url:
        parser.error("Set DATABASE_URL in .env or pass --database-url")
    output = ROOT / "data/private/eval" / f"{args.command}_compare_{args.split}.xlsx"
    if args.command == "scope":
        eval_scope(args.database_url, args.split, args.limit, output)
    elif args.command == "temporal":
        eval_temporal(args.database_url, args.split, args.limit, output)
    elif args.command == "retrieval":
        if args.mode != "bm25":
            suffix = args.mode + (
                f"_{args.embedding_model.split('/')[-1]}" if args.embedding_model else ""
            )
            output = output.with_stem(f"{output.stem}_{suffix}")
        eval_retrieval(
            args.database_url, args.split, args.limit, output, args.mode, args.embedding_model
        )
    elif args.command == "citation":
        eval_citation(args.database_url, args.split, args.limit, output)
    elif args.command == "handoff":
        eval_handoff(args.database_url, args.split, args.limit, args.pin_local_knowledge)


if __name__ == "__main__":
    main()
