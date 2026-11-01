"""Read only the local knowledge sheets; never consume reference findings/observations."""

import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from legal_checker.modules.compliance.criterion_reasoner import SUPPORTED_TYPES
from legal_checker.modules.compliance.decision_rules import parse_json_object
from legal_checker.modules.compliance.numeric_rules import NumericInputError, numeric_kind
from openpyxl import load_workbook


def main():
    root = Path(__file__).resolve().parents[2]
    source = root / (
        "data/private/v750_source_pdf360_r2_supabase/"
        "bang_du_lieu_runtime_v7_5_0_360_STRUCTURED_R2_SUPABASE_DEV_VALIDATION.xlsx"
    )
    with source.open("rb") as stream:
        source_hash = hashlib.file_digest(stream, "sha256").hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)

    def records(name):
        rows = workbook[name].iter_rows(values_only=True)
        columns = next(rows)
        return [dict(zip(columns, row, strict=True)) for row in rows if row and row[0] is not None]

    try:
        criteria = records("tieu_chi_danh_gia")
        rules = records("quy_tac_quyet_dinh")
    finally:
        workbook.close()
    types = Counter(row.get("loai_kiem_tra") for row in criteria)
    shapes = defaultdict(lambda: {"rule_ids": [], "criterion_ids": set(), "rule": None})
    problems = []
    for row in rules:
        try:
            numeric = parse_json_object(row["dieu_kien_json"]).get("numeric_rule")
            if numeric is None:
                continue
            key = json.dumps(numeric, sort_keys=True, ensure_ascii=False, default=str)
            shape = shapes[key]
            shape["rule"] = numeric
            shape["rule_ids"].append(row["quy_tac_id"])
            shape["criterion_ids"].add(row["tieu_chi_id"])
            try:
                shape["engine_kind"] = numeric_kind(numeric)
                shape["supported"] = True
            except NumericInputError as exc:
                shape["supported"] = False
                shape["unknown_reason"] = str(exc)
        except (ValueError, TypeError, KeyError) as exc:
            problems.append({"rule_id": row.get("quy_tac_id"), "error": str(exc)})
    numeric_shapes = []
    for shape in shapes.values():
        shape["criterion_ids"] = sorted(shape["criterion_ids"])
        numeric_shapes.append(shape)
    report = {
        "source": str(source.relative_to(root)),
        "source_workbook_sha256_not_knowledge_snapshot": source_hash,
        "sheets_read": ["tieu_chi_danh_gia", "quy_tac_quyet_dinh"],
        "reference_outputs_or_observations_read": False,
        "criterion_count": len(criteria),
        "decision_rule_count": len(rules),
        "criterion_type_counts": dict(sorted(types.items())),
        "unsupported_types_return_unknown": sorted(set(types) - SUPPORTED_TYPES),
        "numeric_schema_count": len(numeric_shapes),
        "numeric_schemas": numeric_shapes,
        "invalid_policy_rows": problems,
        "interpretation": (
            "Schema support is structural only. Facts, units, denominator, scope and target "
            "must independently pass backend verification; source knowledge is not approved."
        ),
    }
    destination = Path(__file__).with_name("knowledge_inventory.json")
    destination.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {key: value for key, value in report.items() if key != "numeric_schemas"},
            ensure_ascii=True,
        )
    )


if __name__ == "__main__":
    main()
